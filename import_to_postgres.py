"""
One-time migration script: imports from the original SQLite database (shop_mock.db)
and all Excel asset registries into PostgreSQL.

NOTE: shop_mock.db is no longer included in the repository. The live database is
now the authoritative source. To restore from a backup, use maintdb_dump.sql instead:
    psql -d maintdb -f maintdb_dump.sql

Only run this script if you have the original shop_mock.db AND the per-asset
Excel registries under Maintenance_DB/, and need to re-seed from scratch.
Set PG_* env vars if your connection differs from the defaults below.
"""

import os
import sqlite3
from pathlib import Path
from datetime import datetime, timezone

import psycopg2
import psycopg2.extras
import openpyxl

# ── Connection ────────────────────────────────────────────────────────────────
PG_HOST = os.getenv("PG_HOST", "localhost")
PG_PORT = os.getenv("PG_PORT", "5432")
PG_DB   = os.getenv("PG_DB",   "maintdb")
PG_USER = os.getenv("PG_USER", os.getenv("USER", "postgres"))
PG_PASS = os.getenv("PG_PASS", "")

BASE_DIR   = Path(__file__).parent
SQLITE_DB  = BASE_DIR / "shop_mock.db"
MAINT_DIR  = BASE_DIR / "Maintenance_DB"


# ── Helpers ───────────────────────────────────────────────────────────────────

def parse_ts(value):
    """Convert SQLite datetime string or Excel datetime to a tz-aware datetime."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value
    try:
        return datetime.fromisoformat(str(value)).replace(tzinfo=timezone.utc)
    except ValueError:
        return None


def parse_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    try:
        return datetime.fromisoformat(str(value)).date()
    except ValueError:
        return None


def read_field(ws, field_name):
    """Read a value-keyed row from an Excel worksheet (col A = key, col B = value)."""
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == field_name:
            return row[1] if len(row) > 1 else None
    return None


# ── SQLite import ─────────────────────────────────────────────────────────────

def import_sqlite(cur):
    print("Importing SQLite data...")
    src = sqlite3.connect(SQLITE_DB)
    src.row_factory = sqlite3.Row

    # Work Centers — add E001 / Q001 if they appear in labor tickets but not here
    rows = src.execute("SELECT MachineID, MachineName, Department FROM WorkCenters").fetchall()
    extras = [
        ("E001", "Saw / Equipment", "Support"),
        ("Q001", "QC Station",      "QC"),
    ]
    existing_ids = {r["MachineID"] for r in rows}
    for mid, mname, dept in extras:
        if mid not in existing_ids:
            rows.append({"MachineID": mid, "MachineName": mname, "Department": dept})

    psycopg2.extras.execute_batch(cur, """
        INSERT INTO work_centers (machine_id, machine_name, department)
        VALUES (%s, %s, %s)
        ON CONFLICT (machine_id) DO NOTHING
    """, [(r["MachineID"], r["MachineName"], r["Department"]) for r in rows])
    print(f"  work_centers: {len(rows)} rows")

    # Employees
    rows = src.execute("SELECT EmployeeID, Name, Role FROM Employees").fetchall()
    psycopg2.extras.execute_batch(cur, """
        INSERT INTO employees (employee_id, name, role)
        VALUES (%s, %s, %s)
        ON CONFLICT (employee_id) DO NOTHING
    """, [(r["EmployeeID"], r["Name"], r["Role"]) for r in rows])
    print(f"  employees:    {len(rows)} rows")

    # Jobs
    rows = src.execute("""
        SELECT JobNo, PartNo, Customer, OrderQty, DueDate,
               EstSetupHours, EstRunHours, Status
        FROM Jobs
    """).fetchall()
    psycopg2.extras.execute_batch(cur, """
        INSERT INTO jobs (job_no, part_no, customer, order_qty, due_date,
                          est_setup_hours, est_run_hours, status)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (job_no) DO NOTHING
    """, [(r["JobNo"], r["PartNo"], r["Customer"], r["OrderQty"],
           parse_date(r["DueDate"]), r["EstSetupHours"],
           r["EstRunHours"], r["Status"]) for r in rows])
    print(f"  jobs:         {len(rows)} rows")

    # Labor Tickets
    rows = src.execute("""
        SELECT TicketID, EmployeeID, JobNo, MachineID, OpCode, OpDesc,
               TimeIn, TimeOut, SetupHours, RunHours, QtyGood, QtyScrap
        FROM LaborTickets
    """).fetchall()
    psycopg2.extras.execute_batch(cur, """
        INSERT INTO labor_tickets
            (ticket_id, employee_id, job_no, machine_id, op_code, op_desc,
             time_in, time_out, setup_hours, run_hours, qty_good, qty_scrap)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (ticket_id) DO NOTHING
    """, [(r["TicketID"], r["EmployeeID"], r["JobNo"], r["MachineID"],
           r["OpCode"], r["OpDesc"],
           parse_ts(r["TimeIn"]), parse_ts(r["TimeOut"]),
           r["SetupHours"], r["RunHours"], r["QtyGood"], r["QtyScrap"])
          for r in rows])
    print(f"  labor_tickets:{len(rows)} rows")

    src.close()


# ── Asset Registry import ─────────────────────────────────────────────────────

def import_asset(cur, folder: Path):
    xlsx_files = list(folder.glob("Asset_Registry_*.xlsx"))
    if not xlsx_files:
        return

    wb = openpyxl.load_workbook(xlsx_files[0], data_only=True)
    ws_id    = wb["Asset_Identity"]
    ws_admin = wb["Asset_Admin_Metadata"] if "Asset_Admin_Metadata" in wb.sheetnames else None
    ws_notes = wb["Asset_Notes"]          if "Asset_Notes"          in wb.sheetnames else None

    def f(name):
        return read_field(ws_id, name)

    asset_id = f("Asset_ID")
    if not asset_id:
        return

    cur.execute("""
        INSERT INTO assets (
            asset_id, asset_code, department_code, sequence_number,
            display_name, short_name, equipment_category, primary_function,
            manufacturer, model, serial_number, control_type, build_year,
            distributor, distributor_email, distributor_phone, vendor_contacts,
            facility, area_cell, commission_date, asset_status,
            planned_pm_interval_days,
            installation_notes, known_quirks, special_tooling, reference_document,
            template_version, registry_created_by, registry_creation_date,
            final_lock_date, approved_by, folder_path
        ) VALUES (
            %s,%s,%s,%s, %s,%s,%s,%s, %s,%s,%s,%s,%s,
            %s,%s,%s,%s, %s,%s,%s,%s,%s,
            %s,%s,%s,%s, %s,%s,%s,%s,%s,%s
        )
        ON CONFLICT (asset_id) DO NOTHING
    """, (
        asset_id,
        f("Asset_Code"),
        f("Department_Code"),
        f("Asset_Sequence_Number"),
        f("Asset_Display_Name"),
        f("Short_Name"),
        f("Equipment_Category"),
        f("Primary_Function"),
        f("Manufacturer"),
        f("Model"),
        f("Serial_Number"),
        f("Control_Type"),
        f("Build_Year"),
        f("Distributor"),
        f("Distributor_Email"),
        f("Distributor_Phone"),
        f("Vendor_Contacts"),
        f("Facility"),
        f("Area_Cell"),
        parse_date(f("Commission_Date")),
        f("Asset_Status"),
        f("Planned_PM_Interval_Days"),
        read_field(ws_notes, "Installation_Notes") if ws_notes else None,
        read_field(ws_notes, "Known_Quirks")       if ws_notes else None,
        read_field(ws_notes, "Special_Tooling")    if ws_notes else None,
        read_field(ws_notes, "Reference_Document") if ws_notes else None,
        read_field(ws_admin, "Template_Version")         if ws_admin else None,
        read_field(ws_admin, "Registry_Created_By")      if ws_admin else None,
        parse_ts(read_field(ws_admin, "Registry_Creation_Date")) if ws_admin else None,
        parse_date(read_field(ws_admin, "Final_Lock_Date"))      if ws_admin else None,
        read_field(ws_admin, "Approved_By")              if ws_admin else None,
        str(folder),
    ))

    # Batteries
    for n in range(1, 4):
        desc     = f(f"Battery_{n}_Description")
        part_no  = f(f"Battery_{n}_Part_Number")
        interval = f(f"Battery_{n}_Replacement_Interval_Days")
        if desc or part_no:
            cur.execute("""
                INSERT INTO asset_batteries
                    (asset_id, battery_number, description, part_number, replacement_interval_days)
                VALUES (%s, %s, %s, %s, %s)
            """, (asset_id, n, desc, part_no, interval))

    # Peripheral equipment
    for n in range(1, 4):
        name   = f(f"Additional_Equipment_{n}_Name")
        make   = f(f"Additional_Equipment_{n}_Make")
        model  = f(f"Additional_Equipment_{n}_Model")
        serial = f(f"Additional_Equipment_{n}_Serial_Number")
        if name:
            cur.execute("""
                INSERT INTO asset_peripheral_equipment
                    (asset_id, equip_number, name, make, model, serial_number)
                VALUES (%s, %s, %s, %s, %s, %s)
            """, (asset_id, n, name, make, model, serial))


def import_assets(cur):
    print("Importing asset registries...")
    count = 0
    for folder in sorted(MAINT_DIR.iterdir()):
        if folder.is_dir() and not folder.name.startswith("_"):
            import_asset(cur, folder)
            count += 1
    print(f"  assets: {count} folders processed")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    conn = psycopg2.connect(
        host=PG_HOST, port=PG_PORT, dbname=PG_DB,
        user=PG_USER, password=PG_PASS
    )
    conn.autocommit = False
    cur = conn.cursor()

    try:
        import_sqlite(cur)
        import_assets(cur)
        conn.commit()
        print("\nDone. All data committed.")
    except Exception as e:
        conn.rollback()
        print(f"\nERROR: {e}")
        raise
    finally:
        cur.close()
        conn.close()


if __name__ == "__main__":
    main()
