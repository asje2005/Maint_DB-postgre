# Unused after the Postgres migration: app.py no longer imports these.
# Kept only as reference for the old Excel asset-registry workbooks.
from typing import Any
from openpyxl.worksheet.worksheet import Worksheet


def find_field_row(sheet: Worksheet, field_name: str) -> int:
    """Returns the row number where column A equals field_name. Raises ValueError if not found."""
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value == field_name:
            return row
    raise ValueError(f"Required field '{field_name}' not found in sheet '{sheet.title}'")


def write_field(sheet: Worksheet, field_name: str, value: Any) -> None:
    """Writes value to column B of the row where column A equals field_name."""
    row = find_field_row(sheet, field_name)
    sheet.cell(row=row, column=2).value = value


def read_field(sheet: Worksheet, field_name: str) -> Any:
    """Reads column B of the row where column A equals field_name. Returns None if not found."""
    try:
        row = find_field_row(sheet, field_name)
        return sheet.cell(row=row, column=2).value
    except ValueError:
        return None
