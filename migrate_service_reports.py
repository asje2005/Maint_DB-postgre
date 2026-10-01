"""
Import existing Excel service records from Maintenance_DB/ into Postgres.
WIP drafts (Maintenance_DB/_WIP/S-DRAFT-*.xlsx) load as Draft rows;
finalized reports (Maintenance_DB/A-*/Service_Reports/S-*.xlsx) load as Final rows.
Idempotent: re-running skips records already present.

Importable: call ``migrate(engine, db_dir)`` (used by the desktop app). Run as a
script to use the PG_* env vars (defaults to localhost:5432, db maintdb):
    .venv/bin/python migrate_service_reports.py
"""
import os
from pathlib import Path
from datetime import datetime, timezone

from sqlalchemy import create_engine, text

BASE_DIR = Path(__file__).parent
DB_DIR   = BASE_DIR / "Maintenance_DB"

DDL = """
CREATE TABLE IF NOT EXISTS service_reports (
    report_id           SERIAL PRIMARY KEY,
    service_id          TEXT UNIQUE,
    asset_id            TEXT REFERENCES assets(asset_id) ON DELETE SET NULL,
    asset_code          TEXT NOT NULL,
    record_state        TEXT NOT NULL DEFAULT 'Draft',
    service_type        TEXT,
    issue_category      TEXT,
    root_cause          TEXT,
    problem_description TEXT,
    work_performed      TEXT,
    downtime_hours      NUMERIC(6,2),
    return_to_service   TEXT,
    verified_by         TEXT,
    entry_timestamp     TIMESTAMPTZ DEFAULT NOW(),
    finalized_at        TIMESTAMPTZ,
    created_at          TIMESTAMPTZ DEFAULT NOW(),
    updated_at          TIMESTAMPTZ DEFAULT NOW()
);
CREATE INDEX IF NOT EXISTS idx_sr_asset_id     ON service_reports(asset_id);
CREATE INDEX IF NOT EXISTS idx_sr_record_state ON service_reports(record_state);
CREATE INDEX IF NOT EXISTS idx_sr_asset_state  ON service_reports(asset_id, record_state);
"""


def _try_read_excel(path):
    """Return a dict of field values, or empty dict if openpyxl is unavailable."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)

        def rf(sheet_name, field):
            try:
                ws = wb[sheet_name]
                for row in range(1, ws.max_row + 1):
                    if ws.cell(row=row, column=1).value == field:
                        return ws.cell(row=row, column=2).value
            except Exception:
                pass
            return None

        result = {
            "service_type":        rf("Service_Header",        "Service_Type"),
            "problem_description": rf("Problem_and_Work",      "Problem_Description"),
            "work_performed":      rf("Problem_and_Work",      "Work_Performed"),
            "downtime_hours":      rf("Impact_and_Verification","Downtime_Hours"),
            "return_to_service":   rf("Impact_and_Verification","Return_To_Service"),
            "verified_by":         rf("Impact_and_Verification","Verified_By"),
            # The real authored time lives in the workbook; prefer it over the file
            # mtime (which only reflects when the file was last copied/touched).
            "entry_timestamp":     rf("Service_Header",        "Entry_Timestamp"),
        }
        wb.close()
        return result
    except ImportError:
        return {}


def migrate(engine, db_dir=None, verbose=True):
    """Create service_reports (if missing) and import on-disk Excel drafts/finals
    from ``db_dir`` (default Maintenance_DB/). Returns (drafts_imported, finals_imported)."""
    db_dir = Path(db_dir) if db_dir else DB_DIR
    wip_dir = db_dir / "_WIP"

    def log(msg):
        if verbose:
            print(msg)

    with engine.connect() as conn:
        conn.execute(text(DDL))
        conn.commit()

    drafts_imported = 0
    if wip_dir.exists():
        for draft_file in wip_dir.glob("S-DRAFT-*.xlsx"):
            asset_code = draft_file.stem.replace("S-DRAFT-", "")
            asset_id   = f"A-{asset_code}"
            mtime      = datetime.fromtimestamp(draft_file.stat().st_mtime, tz=timezone.utc)
            with engine.connect() as conn:
                exists = conn.execute(
                    text("SELECT 1 FROM service_reports WHERE asset_code = :code AND record_state = 'Draft'"),
                    {"code": asset_code}).fetchone()
            if exists:
                log(f"  Skipping draft {draft_file.name} (already in DB)")
                continue
            data = _try_read_excel(draft_file)
            with engine.connect() as conn:
                conn.execute(text("""
                    INSERT INTO service_reports
                        (asset_id, asset_code, record_state, service_type,
                         problem_description, work_performed, downtime_hours,
                         return_to_service, verified_by, entry_timestamp, created_at, updated_at)
                    VALUES (:aid, :code, 'Draft', :stype, :prob, :work, :dt,
                            :rts, :vb, :ets, :ts, :ts)
                """), {
                    "aid": asset_id, "code": asset_code, "stype": data.get("service_type"),
                    "prob": data.get("problem_description"), "work": data.get("work_performed"),
                    "dt": data.get("downtime_hours"), "rts": data.get("return_to_service"),
                    "vb": data.get("verified_by"), "ets": data.get("entry_timestamp") or mtime, "ts": mtime,
                })
                conn.commit()
            log(f"  Imported draft: {draft_file.name}")
            drafts_imported += 1

    finals_imported = 0
    for asset_dir in db_dir.glob("A-*"):
        if not asset_dir.is_dir() or asset_dir.name == "_WIP":
            continue
        service_dir = asset_dir / "Service_Reports"
        if not service_dir.exists():
            continue
        parts = asset_dir.name.split("-")
        if len(parts) < 2:
            continue
        asset_code = parts[1]
        asset_id   = f"A-{asset_code}"
        for rep_file in service_dir.glob("S-*.xlsx"):
            service_id = rep_file.stem
            mtime      = datetime.fromtimestamp(rep_file.stat().st_mtime, tz=timezone.utc)
            with engine.connect() as conn:
                exists = conn.execute(text("SELECT 1 FROM service_reports WHERE service_id = :sid"),
                                      {"sid": service_id}).fetchone()
            if exists:
                continue
            data = _try_read_excel(rep_file)
            with engine.connect() as conn:
                conn.execute(text("""
                    INSERT INTO service_reports
                        (service_id, asset_id, asset_code, record_state, service_type,
                         problem_description, work_performed, downtime_hours,
                         return_to_service, verified_by, entry_timestamp, finalized_at, created_at, updated_at)
                    VALUES (:sid, :aid, :code, 'Final', :stype, :prob, :work, :dt,
                            :rts, :vb, :ets, :ts, :ts, :ts)
                """), {
                    "sid": service_id, "aid": asset_id, "code": asset_code, "stype": data.get("service_type"),
                    "prob": data.get("problem_description"), "work": data.get("work_performed"),
                    "dt": data.get("downtime_hours"), "rts": data.get("return_to_service"),
                    "vb": data.get("verified_by"), "ets": data.get("entry_timestamp") or mtime, "ts": mtime,
                })
                conn.commit()
            log(f"  Imported finalized: {service_id}")
            finals_imported += 1

    return drafts_imported, finals_imported


def _engine_from_env():
    host = os.getenv("PG_HOST", "localhost")
    port = os.getenv("PG_PORT", "5432")
    db   = os.getenv("PG_DB",   "maintdb")
    user = os.getenv("PG_USER", os.getenv("USER", "postgres"))
    pwd  = os.getenv("PG_PASS", "")
    return create_engine(f"postgresql+psycopg2://{user}:{pwd}@{host}:{port}/{db}")


if __name__ == "__main__":
    drafts, finals = migrate(_engine_from_env())
    print(f"\nMigration complete. Drafts imported: {drafts} | Finalized imported: {finals}")
